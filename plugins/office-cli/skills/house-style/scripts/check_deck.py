#!/usr/bin/env python3
"""Content and craft checks for a house-style .pptx.

    python3 <house-style>/scripts/check_deck.py deck.pptx
    python3 scripts/check_deck.py deck.pptx --palette 'SKS Blue'
    python3 scripts/check_deck.py deck.pptx --template     # a template file, not a deliverable

`officecli view issues` catches overflow and stale fields. `contrast.py` catches
colour that fails the gate. Neither reads what the slide says. This does: the
title column, the agreement between a number in a title and the body under it,
the tells of a deck assembled from a template rather than written.

Exit code 1 if any FAIL. WARN never fails the run — read it and decide.

Craft rules adapted from consulting-pptx-skill (Carnot AI Inc., MIT):
https://github.com/gozen3ji/consulting-pptx-skill — a canon accumulated one line
at a time from real review feedback. The checks below are re-implemented for this
system's canvas, furniture, palettes and bilingual EN·中文 convention.
"""
import argparse
import os
import re
import sys
import unicodedata

import house_prose

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError:
    sys.exit("needs python-pptx:  pip3 install python-pptx --break-system-packages")

EMU_PT = 12700.0

# ── house-style furniture, in pt on the 1440 x 810 canvas ────────────────────
TITLE_Y, TITLE_BAND = 76, 60          # title sits at y=76
RULE_Y = 142                           # title rule
FOOTER_Y = 772                         # footer / page number
CONTENT_BAND = (172.0, 760.0)          # body lives between these
TITLE_W, TITLE_PT = 1328.0, 30.0       # title box width, size
# these five carry their own furniture; the standing-furniture and fill checks skip them
EXEMPT_LAYOUTS = {"cover", "section divider", "closing", "quote", "blank"}

# counters that make a number in a title a promise about the body
COUNTERS = ["段階", "階段", "步驟", "面向", "支柱", "大", "項", "個", "類", "層",
            "steps", "step", "phases", "phase", "pillars", "pillar",
            "stages", "stage", "themes", "theme", "priorities", "areas", "lines"]
NUM_WORD = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6,
            "seven": 7, "eight": 8, "nine": 9, "ten": 10,
            "一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6,
            "七": 7, "八": 8, "九": 9, "十": 10}

fails, warns = [], []
def FAIL(m): fails.append(m)
def WARN(m): warns.append(m)


def hw(text):
    """Half-width length: a CJK/fullwidth glyph counts 2, everything else 1."""
    n = 0
    for ch in text:
        n += 2 if unicodedata.east_asian_width(ch) in ("W", "F", "A") else 1
    return n


def pt(v):
    return (v or 0) / EMU_PT


def shape_text(sh):
    if not sh.has_text_frame:
        return ""
    return "\n".join(p.text for p in sh.text_frame.paragraphs).strip()


def classify(slide, sw, sh_h):
    """Return (title_shape, eyebrow_shape, footer_shapes, body_shapes).

    A real title placeholder outranks geometry. `officecli add slide --prop
    title=…` stamps stock PowerPoint geometry over the layout's slot, and a title
    recognised only by where it sits would disappear from this gate at exactly the
    moment it is in the wrong place — the slide would read as having no title at
    all instead of a misplaced one.
    """
    title = eyebrow = ph_title = None
    footers, body = [], []
    for shp in slide.shapes:
        y, h = pt(shp.top), pt(shp.height)
        txt = shape_text(shp)
        if txt and shp.is_placeholder and "TITLE" in str(shp.placeholder_format.type):
            ph_title = shp
        elif y >= FOOTER_Y - 12:
            footers.append(shp)
        elif abs(y - 48) < 14 and txt:
            eyebrow = shp
        elif abs(y - TITLE_Y) < 20 and txt:
            if title is None or pt(shp.width) > pt(title.width):
                title = shp
        elif y >= CONTENT_BAND[0] - 20:
            body.append(shp)
    return (ph_title or title), eyebrow, footers, body


def title_lines(text):
    """Lines the title will take in a 1328pt box at 30pt bold Arial."""
    per_line = TITLE_W / (TITLE_PT * 0.52)     # ~0.52em average advance, bold
    lines, cur = 1, 0
    for seg in text.split("\n"):
        cur = 0
        for ch in seg:
            w = 2 if unicodedata.east_asian_width(ch) in ("W", "F", "A") else 1
            if cur + w > per_line:
                lines += 1
                cur = 0
            cur += w
        lines += 0
    return lines, cur, per_line


def check_title(idx, text):
    if not text:
        return
    lines, last, per = title_lines(text)
    if lines > 2:
        FAIL(f"p{idx}: title needs {lines} lines in the 1328pt box — cut it, do not shrink the type")
    if lines == 2 and 0 < last <= 3:
        WARN(f"p{idx}: title's second line is {last} half-widths — an orphan; break at the sense join")
    if re.match(r"^\s*[^:：\n]{1,14}[:：]\s*\S", text):
        head = re.split(r"[:：]", text, 1)[0].strip()
        WARN(f"p{idx}: title reads as a label — '{head}: …'. Write the conclusion, not the section name")
    if re.search(r"[。.]\s*$", text):
        WARN(f"p{idx}: title ends in a full stop — titles are message lines, not sentences to close")


def title_number(text):
    m = re.search(r"(\d+|[一二三四五六七八九十]|one|two|three|four|five|six|seven|eight|nine|ten)\s*"
                  r"(" + "|".join(map(re.escape, COUNTERS)) + r")", text, re.I)
    if not m:
        return None
    tok = m.group(1)
    return int(tok) if tok.isdigit() else NUM_WORD.get(tok.lower())


def body_item_count(body_shapes):
    """Count leading enumerators across body text — 01 / 1. / ① / Step 2."""
    seen = set()
    for shp in body_shapes:
        for line in shape_text(shp).split("\n"):
            line = line.strip()
            m = (re.match(r"^(\d{1,2})[.)、．]", line)
            or re.match(r"^0(\d)\b", line)
            or re.match(r"^[Ss]tep\s*(\d)", line)
            or re.match(r"^第\s*([一二三四五六七八九十\d])\s*[階步項]", line))
            if m:
                tok = m.group(1)
                seen.add(int(tok) if tok.isdigit() else NUM_WORD.get(tok, 0))
            elif line[:1] in "①②③④⑤⑥⑦⑧⑨⑩":
                seen.add("①②③④⑤⑥⑦⑧⑨⑩".index(line[0]) + 1)
    seen.discard(0)
    return len(seen)


def signature(text):
    """Crude sentence-shape signature, for spotting a deck written from one mould."""
    t = re.sub(r"[A-Za-z0-9]+", "W", text)
    t = re.sub(r"[一-鿿]+", "C", t)
    return re.sub(r"\s+", "", t)[:24]


def check_fonts(idx, shapes):
    for shp in shapes:
        if not shp.has_text_frame:
            continue
        for para in shp.text_frame.paragraphs:
            for run in para.runs:
                if not run.text.strip():
                    continue
                latin = run.font.name
                if latin and latin != "Arial":
                    FAIL(f"p{idx}: run set in {latin!r} — this system is Arial only")
                # a run with no rPr inherits the master, which is correct. Only a run
                # that sets a latin face and omits the ea face is a real defect: 中文
                # then falls through to whatever the reader's Windows picks.
                if latin and re.search(r"[一-鿿]", run.text):
                    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
                    rpr = run._r.find(A + "rPr")
                    ea = rpr.find(A + "ea") if rpr is not None else None
                    if ea is None or not ea.get("typeface"):
                        FAIL(f"p{idx}: 中文 run sets latin={latin!r} but no font.ea — "
                             f"Windows falls back to a serif")


def check_fill_ratio(idx, body):
    if not body:
        return
    top, bot = CONTENT_BAND
    spans = []
    for shp in body:
        y0, y1 = pt(shp.top), pt(shp.top) + pt(shp.height)
        y0, y1 = max(y0, top), min(y1, bot)
        if y1 > y0:
            spans.append((y0, y1))
    if not spans:
        return
    spans.sort()
    used, cur_s, cur_e = 0.0, *spans[0]
    for s, e in spans[1:]:
        if s > cur_e:
            used += cur_e - cur_s
            cur_s, cur_e = s, e
        else:
            cur_e = max(cur_e, e)
    used += cur_e - cur_s
    ratio = used / (bot - top)
    if ratio < 0.55:
        WARN(f"p{idx}: body fills {ratio:.0%} of the content band — under 55%, the page reads half-empty")


def check_palette(prs, path, palette_name):
    """Every srgbClr on a slide must be in the named palette. Needs contrast.py."""
    here = os.path.dirname(os.path.abspath(__file__))
    for cand in (here,                                              # installed beside contrast.py
                 os.path.join(here, "..", "..", "house-style", "scripts")):
        if os.path.exists(os.path.join(cand, "contrast.py")):
            sys.path.insert(0, os.path.abspath(cand))
            break
    try:
        import contrast
    except Exception as exc:
        WARN(f"palette check skipped — contrast.py not importable ({exc})")
        return
    allowed = contrast.allowed_hexes(palette_name)
    if allowed is None:
        FAIL(f"unknown palette {palette_name!r} — known: {', '.join(contrast.PALETTES)}")
        return
    import zipfile
    z = zipfile.ZipFile(path)
    stray = {}
    for name in z.namelist():
        if not re.match(r"ppt/slides/slide\d+\.xml$", name):
            continue
        idx = int(re.search(r"slide(\d+)", name).group(1))
        for col in set(re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"',
                                  z.read(name).decode("utf-8", "ignore"))):
            if col.upper() not in allowed:
                stray.setdefault(col.upper(), []).append(idx)
    for col, pages in sorted(stray.items()):
        FAIL(f"#{col} is not in {palette_name} — on slide(s) {', '.join(map(str, sorted(pages)))}")


def check_inherited_placeholders(prs):
    """{{…}} that live on a layout rather than a slide.

    In this system the footer, eyebrow and cover furniture sit on the layout, so
    a deck inherits `{{ORG}} | {{DECK_TITLE}}` and renders it on every page while
    the slide's own shapes stay clean — invisible to the per-slide scan, and to
    anyone reading the DOM instead of the render. Only layouts a slide actually
    uses are checked; the other layouts in the template never render.
    """
    # officecli addresses layouts as /slidelayout[N] in presentation order, which is
    # the order python-pptx lists them in — so the FAIL can name the exact path
    # rather than leaving the reader to guess N and clobber the wrong shape.
    nth = {l.part.partname: n for n, l in enumerate(prs.slide_layouts, 1)}
    where, pages = {}, {}
    for i, slide in enumerate(prs.slides, 1):
        lay = slide.slide_layout
        text = "\n".join(s.text_frame.text for s in lay.shapes if s.has_text_frame)
        for ph in set(re.findall(r"\{\{[^}]*\}\}", text)):
            where.setdefault(ph, set()).add(
                f"/slidelayout[{nth.get(lay.part.partname, '?')}] {lay.name or ''}".strip())
            pages.setdefault(ph, []).append(i)
    for ph in sorted(where):
        pp = pages[ph]
        shown = ", ".join(f"p{n}" for n in pp[:8]) + ("…" if len(pp) > 8 else "")
        FAIL(f"{ph} was never replaced — it is on the layout, not the slide, so it "
             f"renders on {shown} while every slide-level check reads clean. It "
             f"lives on {'; '.join(sorted(where[ph]))}. List that layout's shapes "
             f"and set the one that carries it:  officecli get <deck> "
             f"'/slidelayout[N]' --json   then   officecli set <deck> "
             f"'/slidelayout[N]/shape[@id=…]' --prop text=…")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("--palette", help="check every slide colour against this contrast.py palette")
    ap.add_argument("--template", action="store_true",
                    help="a template or layout catalogue: unfilled slides and {{placeholders}} are the point")
    a = ap.parse_args()

    prs = Presentation(a.deck)
    sw, sh_h = pt(prs.slide_width), pt(prs.slide_height)
    if abs(sw - 1440) > 1 or abs(sh_h - 810) > 1:
        FAIL(f"canvas is {sw:.0f} x {sh_h:.0f}pt — this system is 1440 x 810")

    titles, all_text = [], []
    for i, slide in enumerate(prs.slides, 1):
        title_shp, eyebrow, footers, body = classify(slide, sw, sh_h)
        ttl = shape_text(title_shp) if title_shp is not None else ""
        titles.append((i, ttl))
        # (Cover / Divider / Closing titles sit elsewhere on the canvas by design)
        page_text = "\n".join(shape_text(s) for s in slide.shapes)
        all_text.append((i, page_text))

        if not a.template and not slide.shapes:
            FAIL(f"p{i}: no shapes of its own — an empty slide. Every template ships "
                 f"a starter blank slide; remove it once the deck is built:  "
                 f"officecli remove <deck> '/slide[{i}]'")

        if not a.template and "{{" in page_text:
            left = sorted(set(re.findall(r"\{\{[^}]*\}\}", page_text)))
            FAIL(f"p{i}: unreplaced placeholder(s) {' '.join(left)}")

        check_title(i, ttl)
        n = title_number(ttl)
        if n is not None:
            got = body_item_count(body)
            if got and got != n:
                FAIL(f"p{i}: title promises {n} but the body has {got} — "
                     f"the reader sees this first. Fix the number or the body")
        check_fonts(i, slide.shapes)

        layout = (slide.slide_layout.name or "").strip().lower()
        if title_shp is not None and layout not in EXEMPT_LAYOUTS and not a.template:
            ty, tw = pt(title_shp.top), pt(title_shp.width)
            if abs(ty - TITLE_Y) > 20 or abs(tw - TITLE_W) > 24:
                FAIL(f"p{i}: the title box is {tw:.0f}pt wide at y={ty:.0f}pt, not "
                     f"{TITLE_W:.0f}pt at y={TITLE_Y} — it is off the house grid. "
                     f"`add slide --prop title=…` stamps stock geometry over the "
                     f"layout slot; add the placeholder instead:  officecli add "
                     f"<deck> '/slide[{i}]' --type placeholder --prop phType=title")
        if layout not in EXEMPT_LAYOUTS:
            if not a.template:   # a template or catalogue is unfilled on purpose
                check_fill_ratio(i, body)
            # the eyebrow, footer and page number live on the layout in this system,
            # so a slide inherits them. Only flag when neither the slide nor its
            # layout carries them — i.e. someone built on Blank or stripped them.
            l_title, l_eyebrow, l_footers, _ = classify(slide.slide_layout, sw, sh_h)
            if body and eyebrow is None and l_eyebrow is None:
                WARN(f"p{i}: no eyebrow at y=48 on the slide or its layout "
                     f"({slide.slide_layout.name}) — standing furniture is incomplete")
            if body and not footers and not l_footers:
                WARN(f"p{i}: no footer / page number at y=772 on the slide or its layout")

    real = [t for _, t in titles if t]
    if len(real) >= 5:
        sigs = {}
        for t in real:
            sigs.setdefault(signature(t), 0)
            sigs[signature(t)] += 1
        top = max(sigs.values())
        if top / len(real) >= 0.6:
            WARN(f"{top} of {len(real)} titles share one sentence shape — "
                 f"that is a template being filled in, not a deck being written")

    for w in house_prose.scan([t for _, t in all_text], "deck"):
        WARN(w)

    if not a.template:
        if not prs.slides:
            FAIL("the deck has no slides")
        check_inherited_placeholders(prs)

    if a.palette:
        check_palette(prs, a.deck, a.palette)

    print("Title column — read it top to bottom; it should be one argument:\n")
    for i, t in titles:
        print(f"  {i:>3}  {t or '—'}")
    print()
    for w in warns:
        print("WARN ", w)
    for f in fails:
        print("FAIL ", f)
    print(f"\n{len(fails)} FAIL / {len(warns)} WARN")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
