#!/usr/bin/env python3
"""Self-check for the house-style checkers:  python3 scripts/test_checks.py

Builds a deliberately broken .docx and .xlsx, runs each checker over it, and
asserts the defect is reported. A checker that stops catching its own fixture
is worse than no checker, because a green run then means nothing.
"""
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent


def run(script, *args):
    p = subprocess.run([sys.executable, str(HERE / script), *map(str, args)],
                       capture_output=True, text=True)
    return p.returncode, p.stdout + p.stderr


def expect(out, *fragments):
    for f in fragments:
        assert f in out, f"missing {f!r} in:\n{out}"


def docx_fixture(tmp):
    import docx
    d = docx.Document()
    d.add_heading("Board paper {{DECK_TITLE}}", 1)
    d.add_heading("Detail", 3)                       # skips H2
    p = d.add_paragraph("We leverage the 資訊安全 posture, and 資安 too.")
    p.add_run(" wrong face").font.name = "Times New Roman"
    ghost = d.add_paragraph("styled with a style that does not exist")   # silent no-op
    from docx.oxml.ns import qn
    pPr = ghost._p.get_or_add_pPr()
    pPr.append(pPr.makeelement(qn("w:pStyle"), {qn("w:val"): "GhostStyle"}))
    path = tmp / "broken.docx"
    d.save(path)
    return path


def xlsx_fixture(tmp):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item {{UNIT}}", "NT$ thousands"])
    for i in range(7):
        ws.append([f"line {i}", i * 1000])
    ws["A1"].fill = PatternFill("solid", start_color="FFDD0000")
    ws["B2"].number_format = "#,##0;[Red](#,##0)"
    path = tmp / "broken.xlsx"
    wb.save(path)
    return path


def pptx_fixture(tmp):
    """A deck the way one is actually built: the shipped template, one slide added,
    and the starter blank slide left behind. Both defects it carries are inherited
    from the layout, which is exactly the class the per-slide scan cannot see."""
    from pptx import Presentation
    prs = Presentation(HERE.parent / "templates" / "yukima" / "yukima.pptx")
    lay = next(l for l in prs.slide_layouts if (l.name or "") == "Title and Content")
    slide = prs.slides.add_slide(lay)
    if slide.shapes.title is not None:
        slide.shapes.title.text = "Incidents fell 40% after the MFA rollout"
    path = tmp / "deck.pptx"
    prs.save(path)
    return path


def main():
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)

        code, out = run("check_doc.py", docx_fixture(tmp), "--palette", "ANA Blue")
        assert code == 1, f"broken .docx passed:\n{out}"
        expect(out,
               "{{DECK_TITLE}} was never replaced",
               "font 'Times New Roman' is not a house face",
               "heading level jumps H1 → H3",
               "both '資訊安全' and '資安' appear",
               "'leverage'",
               "style 'GhostStyle' is referenced but never defined")

        code, out = run("check_doc.py", docx_fixture(tmp), "--template",
                        "--font", "Times New Roman")
        assert "never replaced" not in out and "not a house face" not in out, out

        code, out = run("check_book.py", xlsx_fixture(tmp), "--palette", "ANA Blue")
        assert code == 1, f"broken .xlsx passed:\n{out}"
        expect(out,
               "{{UNIT}} was never replaced",
               "#DD0000 (fill) is not in the palette",
               "negatives in red",
               "no frozen header")

        code, out = run("check_book.py", xlsx_fixture(tmp), "--template")
        assert "{{UNIT}}" not in out, out

        code, out = run("check_deck.py",
                        HERE.parent / "references" / "layout-catalogue" / "yukima-layouts.pptx",
                        "--template", "--palette", "Yukima")
        assert code == 0, f"the Yukima catalogue should pass its own palette:\n{out}"

        deck = pptx_fixture(tmp)
        code, out = run("check_deck.py", deck, "--palette", "Yukima")
        assert code == 1, f"a deck with an unreplaced footer and a leftover blank slide passed:\n{out}"
        expect(out,
               "{{ORG}} was never replaced",
               "{{DECK_TITLE}} was never replaced",
               "it is on the layout",
               "p1: no shapes of its own")

        code, out = run("check_deck.py", deck, "--template")
        assert code == 0, f"--template should exempt an unfilled deck:\n{out}"

    print("all checks pass")


if __name__ == "__main__":
    main()
