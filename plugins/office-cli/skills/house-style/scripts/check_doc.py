#!/usr/bin/env python3
"""House-style checks for a .docx deliverable.

    python3 <house-style>/scripts/check_doc.py report.docx
    python3 scripts/check_doc.py report.docx --palette 'SKS Blue'
    python3 scripts/check_doc.py form.docx --template      # a template: {{placeholders}} are the point

`officecli view issues` catches overflow and stale fields, `officecli validate`
catches schema breaks. Neither reads what the document says or whether it is on
style. This does: the heading outline, unreplaced placeholders, off-palette
colour, a run set in the wrong face, the East Asian font nobody set, and the
prose tells.

Exit code 1 if any FAIL. WARN never fails the run — read it and decide.
"""
import argparse
import os
import re
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import house_prose

try:
    import docx
except ImportError:
    sys.exit("needs python-docx:  pip3 install python-docx --break-system-packages")

HOUSE_FACES = {"Arial"}          # --font adds to this
CJK = re.compile(r"[㐀-鿿豈-﫿぀-ヿ]")
BODY_PARTS = re.compile(r"^word/(document|header\d*|footer\d*)\.xml$")

fails, warns = [], []
def FAIL(m): fails.append(m)
def WARN(m): warns.append(m)


def parts(path):
    """{part name: xml text} for the parts that carry visible content."""
    z = zipfile.ZipFile(path)
    return {n: z.read(n).decode("utf-8", "ignore")
            for n in z.namelist() if BODY_PARTS.match(n)}


def plain(xml):
    """Visible text of a part, runs joined so a split placeholder still reads whole."""
    body = re.sub(r"<w:p[ >]", "\n<w:p ", xml)
    return re.sub(r"<[^>]+>", "", body)


def check_placeholders(xmls, is_template):
    if is_template:
        return
    for name, xml in xmls.items():
        for ph in sorted(set(re.findall(r"\{\{[^}\n]{0,60}\}\}", plain(xml)))):
            FAIL(f"{ph} was never replaced — in {name}")


def check_palette(xmls, palette_name):
    """Every explicit colour in a visible part must be in the named palette."""
    import contrast
    allowed = contrast.allowed_hexes(palette_name)
    if allowed is None:
        FAIL(f"unknown palette {palette_name!r} — known: {', '.join(contrast.PALETTES)}")
        return
    stray = {}
    for name, xml in xmls.items():
        found = set(re.findall(r'(?:w:val|w:fill|w:color|val)="#?([0-9A-Fa-f]{6})"', xml))
        found |= set(re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', xml))
        for col in found:
            if col.upper() not in allowed:
                stray.setdefault(col.upper(), set()).add(name)
    for col, where in sorted(stray.items()):
        FAIL(f"#{col} is not in {palette_name} — in {', '.join(sorted(where))}")


def check_fonts(xmls, faces):
    """Faces actually named on runs in the visible parts."""
    seen = {}
    for name, xml in xmls.items():
        for attr, face in re.findall(r'w:(ascii|hAnsi|eastAsia|cs)="([^"]+)"', xml):
            if face.startswith("+"):          # theme reference, resolves to the master
                continue
            seen.setdefault(face, set()).add(name)
    for face, where in sorted(seen.items()):
        if face not in faces:
            FAIL(f"font {face!r} is not a house face ({', '.join(sorted(faces))}) — "
                 f"in {', '.join(sorted(where))}")


def check_east_asian(path, xmls):
    """CJK text with no eastAsia font anywhere falls back to the theme default."""
    if not any(CJK.search(plain(x)) for x in xmls.values()):
        return
    styles = zipfile.ZipFile(path)
    pool = "".join(xmls.values())
    for extra in ("word/styles.xml", "word/theme/theme1.xml"):
        try:
            pool += styles.read(extra).decode("utf-8", "ignore")
        except KeyError:
            pass
    if "w:eastAsia" not in pool and "eastAsia" not in pool:
        FAIL("document contains 中文 but no East Asian font is set anywhere — "
             "set font.ea alongside font.latin or it renders in the theme default")


def check_dangling_styles(path, xmls):
    """A style reference the document never defines is a silent no-op.

    `--prop style=Heading1` on a document with no Heading1 definition writes the
    reference and renders as body text. Nothing errors; the heading is simply not
    a heading, and neither is the TOC built from it.
    """
    try:
        styles = zipfile.ZipFile(path).read("word/styles.xml").decode("utf-8", "ignore")
    except KeyError:
        styles = ""
    defined = set(re.findall(r'w:styleId="([^"]+)"', styles))
    used = set()
    for xml in xmls.values():
        used |= set(re.findall(r'w:(?:p|r|tbl)Style w:val="([^"]+)"', xml))
    for name in sorted(used - defined):
        FAIL(f"style {name!r} is referenced but never defined — it renders as body "
             f"text and no TOC will find it; `officecli add <file> /styles --type style` first")


def check_structure(doc):
    """Heading outline, plus the two ways it goes wrong."""
    outline, last = [], 0
    for p in doc.paragraphs:
        m = re.match(r"Heading ?(\d)", p.style.name or "")
        if m and p.text.strip():
            lvl = int(m.group(1))
            outline.append((lvl, p.text.strip()))
            if last and lvl > last + 1:
                WARN(f"heading level jumps H{last} → H{lvl} at {p.text.strip()[:40]!r}")
            last = lvl
    if not outline:
        WARN("no Heading styles used — a report with no outline cannot be navigated "
             "and its TOC will be empty")
    elif outline[0][0] != 1:
        WARN(f"document opens at H{outline[0][0]}, not H1")
    return outline


def check_furniture(doc):
    for i, section in enumerate(doc.sections, 1):
        text = "".join(p.text for p in section.footer.paragraphs).strip()
        if not text:
            WARN(f"section {i} has no footer — house furniture is "
                 "'{{ORG}}  |  {{DECK_TITLE}}' at 9pt")
    for i, table in enumerate(doc.tables, 1):
        if not table.rows:
            continue
        head = table.rows[0]
        if not any(b'w:shd' in c._tc.xml.encode() or 'w:shd' in c._tc.xml
                   for c in head.cells):
            WARN(f"table {i} header row has no fill — house tables head with a "
                 "filled row and white text")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("doc")
    ap.add_argument("--palette", help="check every explicit colour against this contrast.py palette")
    ap.add_argument("--font", action="append", default=[],
                    help="additional permitted font face (repeatable)")
    ap.add_argument("--template", action="store_true",
                    help="a template file, where {{placeholders}} are expected")
    a = ap.parse_args()

    doc = docx.Document(a.doc)
    xmls = parts(a.doc)

    check_placeholders(xmls, a.template)
    check_fonts(xmls, HOUSE_FACES | set(a.font))
    check_east_asian(a.doc, xmls)
    check_dangling_styles(a.doc, xmls)
    outline = check_structure(doc)
    check_furniture(doc)
    if a.palette:
        check_palette(xmls, a.palette)

    texts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        texts += [c.text for r in t.rows for c in r.cells]
    for w in house_prose.scan(texts, "document"):
        WARN(w)

    print("Heading outline — read it top to bottom; it should be the argument:\n")
    for lvl, text in outline or [(1, "—")]:
        print(f"  {'  ' * (lvl - 1)}H{lvl}  {text}")
    print()
    for w in warns:
        print("WARN ", w)
    for f in fails:
        print("FAIL ", f)
    print(f"\n{len(fails)} FAIL / {len(warns)} WARN")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
