#!/usr/bin/env python3
"""House-style checks for an .xlsx deliverable.

    python3 <house-style>/scripts/check_book.py book.xlsx
    python3 scripts/check_book.py book.xlsx --palette 'SKS Blue'
    python3 scripts/check_book.py template.xlsx --template   # {{placeholders}} are the point

`officecli view issues` catches formulas that never evaluated and references to
missing sheets. This catches what survives that: a formula that evaluated to an
error, a header row nobody froze, colour outside the palette, a red negative
where house style says parentheses, and a column too narrow for what is in it.

Exit code 1 if any FAIL. WARN never fails the run — read it and decide.
"""
import argparse
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("needs openpyxl:  pip3 install openpyxl --break-system-packages")

ERRORS = {"#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NULL!", "#NUM!"}
DEFAULT_WIDTH = 8.43
FREEZE_FROM_ROWS = 5          # below this a sheet is a note, not a table

fails, warns = [], []
def FAIL(m): fails.append(m)
def WARN(m): warns.append(m)


def argb(colour):
    """The 6-hex body of an openpyxl colour, or None when it is themed/indexed."""
    if colour is None or colour.type != "rgb" or not isinstance(colour.rgb, str):
        return None
    rgb = colour.rgb.upper()
    return rgb[-6:] if len(rgb) in (6, 8) else None


def cells(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None or c.has_style:
                yield c


def check_sheet(ws, cached, allowed, is_template):
    used = [c for c in cells(ws) if c.value is not None]
    if not used:
        return
    rows = max(c.row for c in used)

    for c in used:
        if isinstance(c.value, str):
            if not is_template:
                for ph in set(re.findall(r"\{\{[^}\n]{0,60}\}\}", c.value)):
                    FAIL(f"{ph} was never replaced — {ws.title}!{c.coordinate}")
            if c.value.strip() in ERRORS:
                FAIL(f"{c.value.strip()} in {ws.title}!{c.coordinate}")
        if "[Red]" in (c.number_format or ""):
            FAIL(f"{ws.title}!{c.coordinate} shows negatives in red — house style is "
                 f"parentheses: '#,##0;(#,##0)'")
        if allowed is not None:
            for hexes, what in ((argb(c.fill.start_color) if c.fill and
                                 c.fill.fill_type else None, "fill"),
                                (argb(c.font.color) if c.font else None, "type")):
                if hexes and hexes not in allowed:
                    FAIL(f"#{hexes} ({what}) is not in the palette — "
                         f"{ws.title}!{c.coordinate}")

    for c in cells(ws):
        val = cached[ws.title][c.coordinate] if ws.title in cached else None
        if isinstance(val, str) and val.strip() in ERRORS:
            FAIL(f"{val.strip()} evaluated in {ws.title}!{c.coordinate}")

    if rows >= FREEZE_FROM_ROWS and not ws.freeze_panes:
        FAIL(f"{ws.title} has {rows} rows and no frozen header — "
             f"set freeze_panes so the header survives scrolling")

    head = [c for c in ws[1] if c.value is not None]
    if rows >= FREEZE_FROM_ROWS and head and not any(c.font and c.font.bold for c in head):
        WARN(f"{ws.title} header row is not bold — house tables head with a filled "
             f"row in white bold type")

    widest = {}
    for c in used if rows >= FREEZE_FROM_ROWS else []:
        col = c.column_letter
        widest[col] = max(widest.get(col, 0), len(str(c.value)))
    for col, need in widest.items():
        have = ws.column_dimensions[col].width or DEFAULT_WIDTH
        if need > have + 1:
            WARN(f"{ws.title} column {col} holds {need} characters in a "
                 f"{have:.0f}-wide column — it will render as ######")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("book")
    ap.add_argument("--palette", help="check every explicit colour against this contrast.py palette")
    ap.add_argument("--template", action="store_true",
                    help="a template file, where {{placeholders}} are expected")
    a = ap.parse_args()

    allowed = None
    if a.palette:
        import contrast
        allowed = contrast.allowed_hexes(a.palette)
        if allowed is None:
            print(f"FAIL  unknown palette {a.palette!r} — "
                  f"known: {', '.join(contrast.PALETTES)}")
            return 1

    wb = load_workbook(a.book)
    evaluated = load_workbook(a.book, data_only=True)
    cached = {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row}
              for ws in evaluated.worksheets}

    print("Sheets — each should be one subject:\n")
    for ws in wb.worksheets:
        print(f"  {ws.title:<24} {ws.max_row:>5} rows x {ws.max_column} cols"
              f"{'' if ws.freeze_panes else '   (no frozen header)'}")
        check_sheet(ws, cached, allowed, a.template)
    print()

    for w in warns:
        print("WARN ", w)
    for f in fails:
        print("FAIL ", f)
    print(f"\n{len(fails)} FAIL / {len(warns)} WARN")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
